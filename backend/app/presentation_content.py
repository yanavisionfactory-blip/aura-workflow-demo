"""Render approved, bounded presentation content as editable widescreen slides."""
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

PHASE_SCHEMA = {"type": "object", "additionalProperties": False,
    "required": ["period", "title", "items"], "properties": {
        "period": {"type": "string", "minLength": 1, "maxLength": 24},
        "title": {"type": "string", "minLength": 1, "maxLength": 40},
        "items": {"type": "array", "minItems": 1, "maxItems": 5,
                  "items": {"type": "string", "minLength": 1, "maxLength": 90}},
        "scene": {"type": "string", "enum": ["rain_window", "paper_boat", "lantern"]}}}
PRESENTATION_SCHEMA = {"type": "object", "additionalProperties": False,
    "required": ["title", "phases"], "properties": {
        "title": {"type": "string", "minLength": 1, "maxLength": 50},
        "subtitle": {"type": "string", "maxLength": 120},
        "layout": {"type": "string", "enum": ["timeline", "slides", "creative"],
                   "description": "Use creative for one populated campaign asset, slides for a deck, or timeline for a roadmap."},
        "canvas": {"type": "string", "enum": ["square", "landscape"],
                   "description": "For creative layouts: square social asset or 16:9 landscape screen asset."},
        "phases": {"type": "array", "minItems": 1, "maxItems": 4, "items": PHASE_SCHEMA}}}


def render_timeline(arguments: dict) -> bytes:
    from jsonschema import validate
    validate(arguments, PRESENTATION_SCHEMA)
    if arguments.get('layout') == 'creative':
        return _render_creative(arguments)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(16), Inches(9)
    def label(slide, x, y, w, h, value, size, color='FFFFFF', bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        p = frame.paragraphs[0]
        p.text = value
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.font.name = 'Aptos'
        p.line_spacing = 1.05
        p.space_after = Pt(0)

    phases = arguments['phases']
    if arguments.get('layout') == 'slides':
        accents = ('64DBC4', '78A9FF', 'F2C14E', 'C69CFF')
        for index, phase in enumerate(phases):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string('101C2C')
            accent = accents[index % len(accents)]
            illustrated = bool(phase.get("scene"))
            if illustrated:
                from .poem_art import render_poem_scene

                slide.shapes.add_picture(BytesIO(render_poem_scene(phase["scene"])),
                                         Inches(8.25), Inches(1.1),
                                         width=Inches(6.85), height=Inches(6.7))
            label(slide, .75, .5, 12.5, .45, arguments['title'], 15, 'B7C7D9', True)
            label(slide, .75, 1.25, 14.4, .45, phase['period'].upper(), 15, accent, True)
            label(slide, .75, 1.85, 14.4, 1.0, phase['title'], 34, 'FFFFFF', True)
            for item_index, item in enumerate(phase['items']):
                label(
                    slide, 1.0, 3.15 + item_index * .92,
                    6.7 if illustrated else 13.8, .8,
                    item if illustrated else f"•  {item}", 19 if illustrated else 18, 'DCE6F1'
                )
            label(slide, .75, 8.25, 13.2, .3, arguments.get('subtitle', ''), 10, '8295AA')
            label(slide, 14.2, 8.25, 1.0, .3, f"{index + 1} / {len(phases)}", 10, accent, True)
    else:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('101C2C')
        label(slide, .7, .5, 14.6, .9, arguments['title'], 34, bold=True)
        label(slide, .7, 1.5, 14.6, .8, arguments.get('subtitle', ''), 18, 'B7C7D9')
        width = 14.6 / len(phases)
        for i, phase in enumerate(phases):
            x = .7 + i * width
            label(slide, x, 2.7, width-.3, .5, phase['period'], 17, '64DBC4', True)
            label(slide, x, 3.4, width-.3, 1.0, phase['title'], 22, bold=True)
            for j, item in enumerate(phase['items']):
                label(slide, x, 4.55+j*.72, width-.35, .7, item, 14, 'DCE6F1')
    result = BytesIO()
    deck.save(result)
    return result.getvalue()


def _render_creative(arguments: dict) -> bytes:
    """One populated, editable campaign asset using the existing import grant.

    Title is the business name; the single phase supplies the disclosure,
    headline, service line and CTA. No external images or live destinations
    are introduced by this layout.
    """
    from pptx.enum.shapes import MSO_SHAPE

    if len(arguments['phases']) != 1 or len(arguments['phases'][0]['items']) != 2:
        raise ValueError('A creative needs one phase with a service line and CTA')
    phase = arguments['phases'][0]
    square = arguments.get('canvas') == 'square'
    deck = Presentation()
    deck.slide_width = Inches(11.25 if square else 20)
    deck.slide_height = Inches(11.25)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('102A36')

    def shape(kind, x, y, width, height, color):
        item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(color)
        item.line.fill.background()
        return item

    def text(x, y, width, height, value, size, color='FAF7F0', bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        p = frame.paragraphs[0]
        p.text = value
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.line_spacing = 1.04
        p.space_after = Pt(0)

    margin = .8 if square else 1.0
    text(margin, .7, 9.5 if square else 17.8, .7, arguments['title'], 26 if square else 32, bold=True)
    text(margin, 1.55, 9.5 if square else 17.8, .4, phase['period'], 13 if square else 16, '67D8C4', True)
    text(margin, 2.7, 7.2 if square else 12, 2.7,
         phase['title'], 59 if square else 76, bold=True)
    text(margin, 6.05, 9.6 if square else 12.2, .95,
         phase['items'][0], 22 if square else 28)
    text(margin, 7.22, 9.6 if square else 12.2, .65,
         arguments.get('subtitle', ''), 15 if square else 19, 'B9CCD1')

    # Editable geometric illustration, not a purported customer property.
    hx, hy, unit = (8.0, 3.0, .65) if square else (14.0, 2.65, 1.35)
    shape(MSO_SHAPE.OVAL, hx-.25*unit, hy+.1*unit, 3.1*unit, 4.2*unit, '183D48')
    shape(MSO_SHAPE.ISOSCELES_TRIANGLE, hx, hy+.2*unit, 2.6*unit, 1.8*unit, '67D8C4')
    shape(MSO_SHAPE.RECTANGLE, hx+.3*unit, hy+1.55*unit, 2.0*unit, 1.85*unit, 'FAF7F0')
    for wx in (.65, 1.5):
        shape(MSO_SHAPE.RECTANGLE, hx+wx*unit, hy+1.98*unit, .5*unit, .55*unit, '102A36')
    for line in range(3):
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, hx+.25*unit+line*.12*unit,
              hy+(3.7+line*.3)*unit, (2.2-line*.3)*unit, .075*unit, '67D8C4')

    cta_width = 9.65 if square else 12.0
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin, 8.45, cta_width, 1.05, '67D8C4')
    text(margin+.28, 8.67, cta_width-.56, .65, phase['items'][1],
         23 if square else 30, '102A36', True)
    text(margin, 10.2, 9.6 if square else 17.8, .45, phase['period'],
         12 if square else 15, 'B9CCD1')
    result = BytesIO()
    deck.save(result)
    return result.getvalue()
