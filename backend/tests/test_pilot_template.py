"""The four-app pilot is bounded by exact input and real provider receipts."""

import asyncio
import json
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from pptx import Presentation

from app import agent_runtime, orchestrator
from app.connection_permissions import verification_permission_fixes
from app.native_connectors import native_manifest
from app.pilot_template import PILOT_OPERATIONS, PILOT_PREFIX, PilotInputError, pilot_template
from app.presentation_content import render_timeline


def inventory(connected=True):
    return [
        {"slug": "google", "connected": connected,
         "allowed_operations": ["docs.create", "docs.get", "calendar.create", "calendar.get",
                                "gmail.send", "gmail.get"]},
        {"slug": "canva", "connected": connected,
         "allowed_operations": ["canva.presentation.create", "canva.import.get",
                                "canva.export.create", "canva.export.get"]},
    ]


def pilot_prompt(**changes):
    start = datetime.now(UTC) + timedelta(days=1)
    fields = {
        "doc_title": "Pilot briefing", "doc_body": "AURA pilot briefing for the team.",
        "event_title": "Pilot review", "event_start": start.isoformat(),
        "event_end": (start + timedelta(minutes=30)).isoformat(),
        "canva_title": "Pilot briefing", "canva_bullets": ["First task", "Second task"],
        "email_to": "me",
        **changes,
    }
    return PILOT_PREFIX + json.dumps(fields)


def test_pilot_compiles_without_a_model_or_inferred_fields(monkeypatch):
    async def no_model(*_args, **_kwargs):
        raise AssertionError("The pilot must not call an AI planner")

    monkeypatch.setattr(orchestrator, "create_plan", no_model)
    connected = inventory()
    plan = asyncio.run(orchestrator._create_compiled_plan(
        pilot_prompt(), connected, set(),
        {item["slug"]: native_manifest(item["slug"]) for item in connected},
    ))
    assert [step.operation for step in plan.steps] == list(PILOT_OPERATIONS)
    event = plan.steps[1]
    assert event.arguments["title"] in event.reason
    assert datetime.fromisoformat(event.arguments["start"]["dateTime"]).strftime(
        "%a %d %b %Y, %I:%M %p UTC%z"
    ) in event.reason
    assert plan.steps[-1].arguments["to"] == "me"
    assert all(step.consequential for step in plan.steps if step.operation != "canva.export.create")
    assert plan.planning_artifacts["planner_recovery_mode"] == "audited_four_app_pilot_v1"
    assert agent_runtime.deterministic_plan_fixes(plan, connected, set()) == []
    assert verification_permission_fixes(plan, connected) == []


@pytest.mark.parametrize("changes", [
    {"email_to": "someone@example.com"},
    {"event_start": "2026-09-24T10:00:00"},
    {"canva_bullets": ["x" * 91]},
    {"doc_body": " "},
    {"canva_title": " "},
])
def test_pilot_rejects_missing_or_unsafe_fields(changes):
    with pytest.raises(PilotInputError):
        pilot_template(pilot_prompt(**changes), inventory())


def test_pilot_prefers_native_account_over_a_second_marketplace_capability():
    duplicate = inventory() + [{"slug": "other", "allowed_operations": ["gmail.send"]}]
    assert pilot_template(pilot_prompt(), duplicate).steps[-1].tool_slug == "google"


def test_pilot_marks_partial_grants_for_reconnection():
    connected = inventory()
    connected[0]["allowed_operations"].remove("calendar.create")
    connected[0]["allowed_operations"].remove("gmail.get")
    plan = pilot_template(pilot_prompt(), connected)
    assert plan.planning_artifacts["connection_requirements"] == ["google"]


POEM = (
    "Morning opens the window\nThe rain turns gold\nA new day waits\nA small hope wakes\n\n"
    "A paper boat takes the river\nIt crosses the silent street\nThe sunlight finds its way\nAnd warms the world again\n\n"
    "Carry a lantern in the night\nHold its warmth against the wind\nShare the light with others\nAnd welcome tomorrow together"
)


def test_illustrated_poem_is_imported_into_canva_and_mailed_with_full_text():
    prompt = pilot_prompt(doc_title="When Tomorrow Opens", doc_body=POEM,
                          illustrated_poem=True)
    plan = pilot_template(prompt, inventory())
    presentation = plan.steps[2].arguments
    assert [phase["scene"] for phase in presentation["phases"]] == [
        "rain_window", "paper_boat", "lantern"]
    deck = Presentation(BytesIO(render_timeline(presentation)))
    assert len(deck.slides) == 3
    for slide, phase in zip(deck.slides, presentation["phases"], strict=True):
        assert len([shape for shape in slide.shapes if shape.shape_type == 13]) == 1
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert all(line in text for line in phase["items"])
    email = plan.steps[-1].arguments
    assert POEM in email["body"]
    assert email["to"] == "me"
    assert email["attachments"][0]["filename"] == "Illustrated poem.pdf"


def test_illustrated_poem_rejects_text_without_three_complete_verses():
    with pytest.raises(PilotInputError, match="three verses"):
        pilot_template(pilot_prompt(illustrated_poem=True), inventory())


def receipts():
    return [
        {"step_id": f"step-{index}", "operation": operation,
         "provider_result": {"id": f"provider-{index}"},
         "critic": {"action": "accept"}, "outcome_check": {"status": "verified"}}
        for index, operation in enumerate(PILOT_OPERATIONS)
    ]


def test_pilot_final_review_uses_exact_verified_receipts_without_a_model(monkeypatch):
    async def no_model(*_args, **_kwargs):
        raise AssertionError("Verified pilot receipts must not require model credits")

    monkeypatch.setattr(agent_runtime, "_run", no_model)
    evidence = receipts()
    delivered = asyncio.run(agent_runtime.synthesize_result(pilot_prompt(), evidence))
    checked = asyncio.run(agent_runtime.verify_outcome(
        pilot_prompt(), {"planning_artifacts": {"planner_recovery_mode": "audited_four_app_pilot_v1"}},
        evidence, delivered.model_dump(mode="json"),
    ))
    assert delivered.validation_passed is True
    assert checked.status == "verified"
    evidence[2]["outcome_check"]["status"] = "unverified"
    assert asyncio.run(agent_runtime.verify_outcome(
        pilot_prompt(), {"planning_artifacts": {"planner_recovery_mode": "audited_four_app_pilot_v1"}},
        evidence,
    )).status == "unverified"
