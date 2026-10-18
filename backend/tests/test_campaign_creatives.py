from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from pptx import Presentation

from app.presentation_content import render_timeline
from app.providers import ProviderExecutor


def creative(canvas):
    return {
        'title': 'AURA Demo Home Comfort', 'layout': 'creative', 'canvas': canvas,
        'subtitle': 'Dallas · Plano · Richardson · Garland · Irving',
        'phases': [{'period': 'FICTIONAL DEMONSTRATION', 'title': 'Practical HVAC Help',
                    'items': ['AC repair · Maintenance · Heat-pump repair',
                              'Request a service consultation']}],
    }


@pytest.mark.parametrize('canvas,ratio', [('square', 1), ('landscape', 16/9)])
def test_creatives_have_correct_canvas_and_complete_editable_copy(canvas, ratio):
    args = creative(canvas)
    deck = Presentation(BytesIO(render_timeline(args)))
    assert len(deck.slides) == 1
    assert deck.slide_width / deck.slide_height == pytest.approx(ratio)
    text = [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    for value in [args['title'], args['subtitle'], args['phases'][0]['period'],
                  args['phases'][0]['title'], *args['phases'][0]['items']]:
        assert value in text
    for shape in deck.slides[0].shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= deck.slide_width
        assert shape.top + shape.height <= deck.slide_height
    assert not any(shape.shape_type == 13 for shape in deck.slides[0].shapes)


def test_creative_rejects_extra_phases_or_missing_cta():
    args = creative('square')
    args['phases'][0]['items'].pop()
    with pytest.raises(ValueError, match='service line and CTA'):
        render_timeline(args)


async def test_png_export_preserves_requested_pixel_dimensions():
    executor = ProviderExecutor({'access_token': 'test'})
    executor._canva_request = AsyncMock(return_value={'job': {'id': 'export-one'}})
    await executor._canva_export_create({'design_id': 'design-one', 'format': 'png',
                                         'width': 1920, 'height': 1080})
    executor._canva_request.assert_awaited_once_with(
        'POST', 'exports', json={'design_id': 'design-one',
                                'format': {'type': 'png', 'width': 1920, 'height': 1080}})
