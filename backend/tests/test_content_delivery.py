import base64
from copy import deepcopy
from email import policy
from email.parser import BytesParser
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
from pptx import Presentation

from app import file_delivery
from app.native_connectors import NativeConnectorError
from app.orchestrator import _include_requested_story_in_email
from app.outcome_checks import build_outcome_check, evaluate_outcome_check
from app.presentation_content import render_timeline
from app.providers import ProviderExecutor
from app.schemas import PlanStep, WorkflowPlan

PDF = b'%PDF-1.7\nroadmap fixture\n%%EOF'
URL = 'https://export-download.canva.com/fixture.pdf'


def test_requested_story_and_pdf_appear_in_the_exact_approved_email():
    story = 'A paper boat followed the lantern across the pond.'
    plan = WorkflowPlan(name='Story delivery', interpretation='Email story and PDF', steps=[
        PlanStep(key='doc', agent='Docs', tool_slug='google', operation='docs.create',
                 arguments={'title': 'Paper Boat', 'body': story}, reason='Write story',
                 expected_output='Document'),
        PlanStep(key='mail', agent='Mail', tool_slug='google', operation='gmail.send',
                 arguments={'to': 'me', 'body': 'The story is attached.',
                            'attachments': [{'filename': 'Illustrations.pdf',
                                             'url': '{{steps.export.job.urls[0]}}'}]},
                 reason='Send story and PDF', expected_output='Delivery', depends_on=['doc']),
    ])
    request = 'Write a story in Google Docs, then email the story and PDF to me.'

    _include_requested_story_in_email(plan, request)
    assert story in plan.steps[1].arguments['body']
    assert 'The story is attached.' not in plan.steps[1].arguments['body']
    assert plan.steps[1].arguments['attachments'][0]['filename'] == 'Illustrations.pdf'
    _include_requested_story_in_email(plan, request)
    assert plan.steps[1].arguments['body'].count(story) == 1

    plan.steps[1].arguments['attachments'] = []
    with pytest.raises(NativeConnectorError, match='PDF attachment'):
        _include_requested_story_in_email(plan, request)


def test_populated_timeline_has_one_slide_and_all_approved_text():
    a = {'title': '90-day roadmap', 'subtitle': 'Product development', 'phases': [
        {'period': 'Days 1–30', 'title': 'Build', 'items': ['Planner and executor', 'Durable execution']},
        {'period': 'Days 31–60', 'title': 'Verify', 'items': ['Connector contracts']},
        {'period': 'Days 61–90', 'title': 'Release', 'items': ['Performance evaluation']}]}
    deck = Presentation(BytesIO(render_timeline(a)))
    assert len(deck.slides) == 1
    texts = [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    assert a['title'] in texts and a['subtitle'] in texts
    for phase in a['phases']:
        assert phase['title'] in texts and phase['period'] in texts
        assert all(item in texts for item in phase['items'])


def test_slide_layout_creates_one_editable_slide_per_phase():
    arguments = {
        'title': 'Munich weather',
        'subtitle': 'Updated now · Source: Open-Meteo',
        'layout': 'slides',
        'phases': [
            {'period': 'Today', 'title': "Today's conditions", 'items': ['Cool and dry']},
            {'period': 'Next 3 days', 'title': 'Forecast', 'items': ['Tuesday', 'Wednesday']},
            {'period': 'Practical guide', 'title': 'What to wear', 'items': ['Dress in layers']},
        ],
    }

    deck = Presentation(BytesIO(render_timeline(arguments)))

    assert len(deck.slides) == 3
    for index, phase in enumerate(arguments['phases']):
        texts = [
            shape.text for shape in deck.slides[index].shapes if shape.has_text_frame
        ]
        assert arguments['title'] in texts
        assert phase['title'] in texts
        assert any(phase['items'][0] in text for text in texts)


@pytest.mark.parametrize('url', ['http://export-download.canva.com/x', 'https://canva.com.evil.test/x',
    'https://127.0.0.1/a', 'https://export-download.canva.com@evil.test/x',
    'https://export-download.canva.com:444/x', 'file:///etc/passwd'])
def test_download_rejects_untrusted_destinations(url):
    with pytest.raises(ValueError):
        file_delivery.allowed_download_url(url)


async def test_pdf_must_come_from_this_runs_verified_export(monkeypatch):
    download = AsyncMock(return_value=PDF)
    monkeypatch.setattr(file_delivery, 'download_pdf', download)
    a = {'to': 'me', 'body': 'Attached', 'attachments': [{'filename': 'roadmap.pdf', 'url': URL}]}
    with pytest.raises(ValueError):
        await file_delivery.prepare_attachments(a, set())
    download.assert_not_awaited()
    ready = await file_delivery.prepare_attachments(a, {URL})
    assert ready['attachments'][0]['sha256'] == file_delivery.fingerprint(PDF)['sha256']
    assert 'sha256' not in a['attachments'][0]


async def test_gmail_sends_real_pdf_and_rejects_changed_file_before_post(monkeypatch):
    monkeypatch.setattr(file_delivery, 'download_pdf', AsyncMock(return_value=PDF))
    executor = ProviderExecutor({'access_token': 'test'})
    executor._request = AsyncMock(return_value={'id': 'sent-id'})
    a = {'to': 'owner@example.test', 'subject': 'Roadmap', 'body': 'Attached',
         'attachments': [{'filename': 'roadmap.pdf', 'url': URL, **file_delivery.fingerprint(PDF)}]}
    receipt = await executor._gmail_send(a)
    sent = executor._request.call_args.kwargs['json']['raw']
    message = BytesParser(policy=policy.default).parsebytes(base64.urlsafe_b64decode(sent+'='*(-len(sent)%4)))
    assert next(iter(message.iter_attachments())).get_payload(decode=True) == PDF
    executor._request.reset_mock()
    monkeypatch.setattr(file_delivery, 'download_pdf', AsyncMock(return_value=PDF+b'changed'))
    with pytest.raises(ValueError):
        await executor._gmail_send(a)
    executor._request.assert_not_awaited()
    observed = {'id': 'sent-id', 'labelIds': ['SENT'], 'payload': {'headers': [
        {'name': 'To', 'value': a['to']}, {'name': 'Subject', 'value': a['subject']}], 'parts': [
            {'mimeType': 'text/plain', 'body': {'data': base64.urlsafe_b64encode(b'Attached').decode()}},
            {'mimeType': 'application/pdf', 'filename': 'roadmap.pdf', 'body': {'data': base64.urlsafe_b64encode(PDF).decode()}}]}}
    check = build_outcome_check('gmail.send', a, receipt)
    assert evaluate_outcome_check(check, observed)['status'] == 'verified'
    observed['payload']['parts'].pop()
    assert evaluate_outcome_check(check, observed)['status'] == 'failed'


async def test_verified_async_job_result_reaches_downstream_without_another_write(monkeypatch):
    from app import orchestrator
    result = {'job': {'id': 'job-1', 'status': 'in_progress'}}
    terminal = {'job': {'id': 'job-1', 'status': 'success', 'urls': [URL]}}
    step = SimpleNamespace(operation='canva.export.create', output={'provider_result': deepcopy(result)})
    monkeypatch.setattr(orchestrator, 'check_provider_outcome', AsyncMock(return_value={'status': 'verified', 'observed': terminal}))
    session = SimpleNamespace(commit=AsyncMock())
    decision = await orchestrator.review_recorded_result(session, None, step, None, {}, result)
    assert decision.action == 'accept'
    assert result['job']['urls'] == [URL] and step.output['provider_result']['job']['status'] == 'success'


def test_canva_import_polls_saved_job_and_checks_design_identity():
    check = build_outcome_check('canva.presentation.create', {'title': 'Roadmap'}, {'job': {'id': 'job-1'}})
    assert evaluate_outcome_check(check, {'job': {'id': 'job-1', 'status': 'in_progress'}})['status'] == 'pending'
    result = {'job': {'id': 'job-1', 'status': 'success', 'result': {'designs': [{'id': 'design-1', 'title': 'Roadmap'}]}}}
    assert evaluate_outcome_check(check, result)['status'] == 'verified'
    result['job']['id'] = 'wrong-job'
    assert evaluate_outcome_check(check, result)['status'] == 'failed'


async def test_download_redirect_cannot_escape_canva_and_html_is_not_a_pdf(monkeypatch):
    import httpx
    real_client = httpx.AsyncClient
    calls = []
    def redirect(request):
        calls.append(request)
        return httpx.Response(302, headers={'location': 'http://127.0.0.1/private'})
    monkeypatch.setattr(httpx, 'AsyncClient', lambda **kw: real_client(**kw, transport=httpx.MockTransport(redirect)))
    with pytest.raises(ValueError):
        await file_delivery.download_pdf(URL)
    assert len(calls) == 1 and 'authorization' not in calls[0].headers
    monkeypatch.setattr(httpx, 'AsyncClient', lambda **kw: real_client(**kw, transport=httpx.MockTransport(lambda req: httpx.Response(200, content=b'<html>expired link</html>'))))
    with pytest.raises(ValueError):
        await file_delivery.download_pdf(URL)


async def test_canva_import_posts_real_populated_presentation_once(monkeypatch):
    import json

    import httpx
    calls = []
    def handler(request):
        calls.append(request)
        assert request.url.path == '/rest/v1/imports'
        assert request.headers['content-type'] == 'application/octet-stream'
        metadata = json.loads(request.headers['import-metadata'])
        assert base64.b64decode(metadata['title_base64']) == b'Roadmap'
        assert len(Presentation(BytesIO(request.content)).slides) == 1
        return httpx.Response(200, json={'job': {'id': 'import-1', 'status': 'in_progress'}})
    real_client = httpx.AsyncClient
    monkeypatch.setattr(httpx, 'AsyncClient', lambda **kw: real_client(**kw, transport=httpx.MockTransport(handler)))
    result = await ProviderExecutor({'access_token': 'private'})._canva_presentation_create(
        {'title': 'Roadmap', 'phases': [{'period': 'Day 1', 'title': 'Build', 'items': ['First milestone']}]})
    assert result['job']['id'] == 'import-1' and len(calls) == 1


async def test_canva_export_posts_exact_design_once_without_metadata_probe():
    executor = ProviderExecutor({'access_token': 'private'})
    executor._canva_request = AsyncMock(return_value={'job': {'id': 'export-1'}})

    result = await executor._canva_export_create({'design_id': 'design-1', 'format': 'pdf'})

    assert result['job']['id'] == 'export-1'
    executor._canva_request.assert_awaited_once_with(
        'POST',
        'exports',
        json={'design_id': 'design-1', 'format': {'type': 'pdf'}},
    )


async def test_canva_export_does_not_replay_an_uncertain_failure():
    request = httpx.Request('POST', 'https://api.canva.com/rest/v1/exports')
    uncertain = httpx.ReadTimeout('provider response lost', request=request)
    executor = ProviderExecutor({'access_token': 'private'})
    executor._canva_request = AsyncMock(side_effect=uncertain)

    with pytest.raises(httpx.ReadTimeout):
        await executor._canva_export_create({'design_id': 'design-1', 'format': 'pdf'})

    executor._canva_request.assert_awaited_once()


async def test_canva_export_surfaces_definitive_rejection_for_durable_recovery():
    request = httpx.Request('POST', 'https://api.canva.com/rest/v1/exports')
    not_ready = httpx.HTTPStatusError(
        'not ready',
        request=request,
        response=httpx.Response(
            404,
            request=request,
            json={'code': 'design_not_found', 'message': 'Design not found'},
        ),
    )
    executor = ProviderExecutor({'access_token': 'private'})
    executor._canva_request = AsyncMock(side_effect=not_ready)

    with pytest.raises(httpx.HTTPStatusError):
        await executor._canva_export_create({'design_id': 'design-1', 'format': 'pdf'})

    executor._canva_request.assert_awaited_once()


async def test_gmail_receipt_reads_attachment_bytes_only_when_requested(monkeypatch):
    executor = ProviderExecutor({'access_token': 'private'})
    payload = {'id': 'm1', 'payload': {'parts': [{'filename': 'roadmap.pdf', 'body': {'attachmentId': 'a1', 'size': len(PDF)}}]}}
    executor._request = AsyncMock(side_effect=[deepcopy(payload), {'data': base64.urlsafe_b64encode(PDF).decode(), 'size': len(PDF)}])
    result = await executor._gmail_get({'message_id': 'm1', 'verify_attachments': True})
    assert file_delivery.gmail_attachment_fingerprints(result['payload'])[0]['sha256'] == file_delivery.fingerprint(PDF)['sha256']
    assert executor._request.call_args.args[1].endswith('/m1/attachments/a1')
